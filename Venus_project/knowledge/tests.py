from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.test import TestCase
from unittest.mock import patch

from users.models import User

from .models import Category, Document, DocumentChunk, DocumentVersion
from .embeddings import (
    EmbeddingProvider,
    StubEmbeddingProvider,
    get_embedding_provider,
    reset_embedding_provider_cache,
)
from .services import (
    DocumentIngestionError,
    EmptyDocumentError,
    MissingDocumentFileError,
    UnsupportedFileTypeError,
    chunk_text,
    extract_document_text,
    ingest_document_version,
    normalize_text,
    WhitespaceTokenCounter,
)

class KnowledgeModelsTests(TestCase):

    def setUp(self):
        self.user = User.objects.create_user(username='demo_user', password='A secure password 123!')
        self.category = Category.objects.create(
            name='DEMO - Administrative',
            slug='demo-administrative',
            description='Données de démonstration.',
        )
        reset_embedding_provider_cache()

    def tearDown(self):
        reset_embedding_provider_cache()
        super().tearDown()

    def test_stub_embedding_provider_embeds_text_with_configured_dimension(self):
        provider = StubEmbeddingProvider(dimension=5)

        embedding = provider.embed('DEMO - texte')

        self.assertEqual(len(embedding), 5)
        self.assertTrue(all(isinstance(value, float) for value in embedding))
        self.assertEqual(provider.provider_name, 'stub')
        self.assertIsNone(provider.model_name)
        self.assertEqual(provider.dimension, 5)

    def test_stub_embedding_provider_embeds_many_texts(self):
        provider = StubEmbeddingProvider(dimension=3)

        embeddings = provider.embed_many(['premier texte', 'second texte'])

        self.assertEqual(len(embeddings), 2)
        self.assertTrue(all(len(embedding) == 3 for embedding in embeddings))

    def test_stub_embedding_provider_is_deterministic(self):
        provider = StubEmbeddingProvider(dimension=4)

        self.assertEqual(provider.embed('texte stable'), provider.embed('texte stable'))
        self.assertEqual(
            provider.embed_many(['a', 'b']),
            [provider.embed('a'), provider.embed('b')],
        )

    def test_stub_embedding_provider_handles_empty_text(self):
        provider = StubEmbeddingProvider(dimension=4)

        self.assertEqual(provider.embed(''), [0.0, 0.0, 0.0, 0.0])

    def test_embedding_provider_protocol_contract(self):
        provider = StubEmbeddingProvider()

        self.assertIsInstance(provider, EmbeddingProvider)

    def test_embedding_provider_configuration_defaults_to_stub(self):
        provider = get_embedding_provider()

        self.assertIsInstance(provider, StubEmbeddingProvider)
        self.assertEqual(provider.dimension, 8)

    def test_embedding_provider_configuration_selects_supported_provider(self):
        from django.test import override_settings

        with override_settings(M2_EMBEDDING_PROVIDER='stub', M2_EMBEDDING_DIMENSION=6):
            reset_embedding_provider_cache()
            provider = get_embedding_provider()

        self.assertIsInstance(provider, StubEmbeddingProvider)
        self.assertEqual(provider.dimension, 6)

    def test_embedding_provider_configuration_rejects_unsupported_provider(self):
        from django.test import override_settings

        with override_settings(M2_EMBEDDING_PROVIDER='unsupported'):
            reset_embedding_provider_cache()

            with self.assertRaises(ValueError):
               get_embedding_provider()

    def test_category_creation_and_uniqueness(self):
        self.assertEqual(Category.objects.count(), 1)
        self.assertEqual(str(self.category), 'DEMO - Administrative')

        with self.assertRaises(Exception):
            Category.objects.create(name='DEMO - Administrative', slug='demo-admin-2')

        with self.assertRaises(Exception):
            Category.objects.create(name='DEMO - New', slug='demo-administrative')

    def test_document_creation_and_relations(self):
        document = Document.objects.create(
            title='DEMO - Document test',
            description='Document de démonstration.',
            category=self.category,
            access_level=Document.AccessLevel.STUDENT,
            status=Document.Status.ACTIVE,
            created_by=self.user,
        )

        self.assertEqual(str(document), 'DEMO - Document test')
        self.assertEqual(document.category, self.category)
        self.assertEqual(document.status, Document.Status.ACTIVE)
        self.assertEqual(document.access_level, Document.AccessLevel.STUDENT)

    def test_document_versions_are_versioned_per_document(self):
        document = Document.objects.create(
            title='DEMO - Versioned doc',
            category=self.category,
            created_by=self.user,
        )

        file_one = SimpleUploadedFile('demo-v1.txt', b'premiere version', content_type='text/plain')
        version_one = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=file_one,
            original_filename='demo-v1.txt',
            uploaded_by=self.user,
        )

        file_two = SimpleUploadedFile('demo-v2.txt', b'deuxieme version', content_type='text/plain')
        version_two = DocumentVersion.objects.create(
            document=document,
            version_number=2,
            file=file_two,
            original_filename='demo-v2.txt',
            uploaded_by=self.user,
        )

        self.assertEqual(document.versions.count(), 2)
        self.assertEqual(version_one.version_number, 1)
        self.assertEqual(version_two.version_number, 2)
        self.assertTrue(version_one.checksum)
        self.assertTrue(version_two.checksum)

    def test_document_version_number_is_unique_per_document(self):
        document = Document.objects.create(
            title='DEMO - Unique version',
            category=self.category,
            created_by=self.user,
        )

        DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('demo.txt', b'abc', content_type='text/plain'),
            original_filename='demo.txt',
            uploaded_by=self.user,
        )

        with self.assertRaises(Exception):
            DocumentVersion.objects.create(
                document=document,
                version_number=1,
                file=SimpleUploadedFile('demo-2.txt', b'def', content_type='text/plain'),
                original_filename='demo-2.txt',
                uploaded_by=self.user,
            )

    def test_document_chunk_creation_and_constraints(self):
        document = Document.objects.create(
            title='DEMO - Chunked doc',
            category=self.category,
            created_by=self.user,
        )
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('demo.txt', b'chunk content', content_type='text/plain'),
            original_filename='demo.txt',
            uploaded_by=self.user,
        )

        chunk_one = DocumentChunk.objects.create(
            document_version=version,
            chunk_index=0,
            content='Premier chunk.',
            page_number=1,
            section='Introduction',
            metadata={'source': 'demo'},
        )
        chunk_two = DocumentChunk.objects.create(
            document_version=version,
            chunk_index=1,
            content='Second chunk.',
            page_number=2,
            section='Règlement',
            metadata={'source': 'demo'},
        )

        self.assertEqual(version.chunks.count(), 2)
        self.assertEqual(str(chunk_one), f'{version} chunk 0')
        self.assertEqual(chunk_two.page_number, 2)
        self.assertEqual(chunk_two.section, 'Règlement')

    def test_document_version_file_stays_local_and_not_on_e_drive(self):
        document = Document.objects.create(
            title='DEMO - Media path',
            category=self.category,
            created_by=self.user,
        )
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('demo.txt', b'content', content_type='text/plain'),
            original_filename='demo.txt',
            uploaded_by=self.user,
        )

        self.assertIn('knowledge/documents', version.file.name)
        self.assertNotIn('E:', version.file.name)
        self.assertNotIn('e:', version.file.name.lower())

    def test_chunk_index_is_unique_for_same_version(self):
        document = Document.objects.create(
            title='DEMO - Duplicate chunk check',
            category=self.category,
            created_by=self.user,
        )
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('demo.txt', b'content', content_type='text/plain'),
            original_filename='demo.txt',
            uploaded_by=self.user,
        )

        DocumentChunk.objects.create(
            document_version=version,
            chunk_index=0,
            content='A',
        )

        with self.assertRaises(Exception):
            DocumentChunk.objects.create(
                document_version=version,
                chunk_index=0,
                content='B',
            )

    def test_normalize_text_removes_excessive_whitespace(self):
        text = '  Bonjour    monde\n\n\t texte  \n\n encore  \n'
        self.assertEqual(normalize_text(text), 'Bonjour monde\ntexte\nencore')

    def test_chunk_text_keeps_order_and_overlap(self):
        text = ' '.join(f'phrase-{i}' for i in range(20))
        chunks = chunk_text(text, chunk_size=5, overlap=1, token_counter=WhitespaceTokenCounter())
        self.assertTrue(chunks)
        self.assertTrue(chunks[0].startswith('phrase-0'))
        self.assertEqual(chunks[0].split()[0], 'phrase-0')
        self.assertTrue(all('phrase-' in chunk for chunk in chunks))
        self.assertTrue(all(len(chunk.split()) <= 5 for chunk in chunks))
        self.assertEqual(chunks[0].split()[-1], chunks[1].split()[0])

    def test_chunk_text_uses_injected_token_counter(self):
        class CharacterTokenCounter:
            def tokenize(self, text):
                return list(text)

            def detokenize(self, tokens):
                return ''.join(tokens)

        chunks = chunk_text('abcdefgh', chunk_size=4, overlap=1, token_counter=CharacterTokenCounter())

        self.assertEqual(chunks, ['abcd', 'defg', 'gh'])

    def test_chunk_text_accepts_non_string_tokens(self):
        class NumericTokenCounter:
            def tokenize(self, text):
                return list(range(len(text.replace(' ', ''))))

            def detokenize(self, tokens):
                return ''.join(str(token) for token in tokens)

        chunks = chunk_text('abcd efgh', chunk_size=4, overlap=1, token_counter=NumericTokenCounter())

        self.assertEqual(chunks, ['0123', '3456', '67'])

    def test_chunk_text_is_deterministic(self):
        text = ' '.join(f'token-{index}' for index in range(12))
        self.assertEqual(chunk_text(text, 4, 1), chunk_text(text, 4, 1))

    def test_extract_document_text_supports_txt(self):
        temp_path = Path('D:/Venus_M2/Venus_project/test_data_demo/demo.txt')
        temp_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path.write_text('Bonjour monde\nDeuxieme ligne\n', encoding='utf-8')
        try:
            pages = extract_document_text(str(temp_path), '.txt')
            self.assertEqual(pages[0]['text'], 'Bonjour monde\nDeuxieme ligne')
        finally:
            temp_path.unlink(missing_ok=True)

    def test_ingestion_creates_chunks_and_is_idempotent(self):
        document = Document.objects.create(
            title='DEMO - Ingestion test',
            category=self.category,
            created_by=self.user,
        )
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('demo.txt', b'hello world ' * 30, content_type='text/plain'),
            original_filename='demo.txt',
            uploaded_by=self.user,
        )

        first_chunks = ingest_document_version(version, chunk_size=30, overlap=5)
        second_chunks = ingest_document_version(version, chunk_size=30, overlap=5)

        self.assertEqual(len(first_chunks), len(second_chunks))
        self.assertTrue(all(chunk.content for chunk in second_chunks))
        self.assertEqual(version.chunks.count(), len(second_chunks))

    def test_ingestion_raises_for_missing_file_and_unknown_extension(self):
        document = Document.objects.create(
            title='DEMO - Error handling',
            category=self.category,
            created_by=self.user,
        )
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('demo.bin', b'abc', content_type='application/octet-stream'),
            original_filename='demo.bin',
            uploaded_by=self.user,
        )

        with self.assertRaises(UnsupportedFileTypeError):
            ingest_document_version(version)

        version.file = 'missing.txt'
        with self.assertRaises(MissingDocumentFileError):
            ingest_document_version(version)

    def test_empty_document_raises_clean_error(self):
        temp_path = Path('D:/Venus_M2/Venus_project/test_data_demo/empty.txt')
        temp_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path.write_text('', encoding='utf-8')
        try:
            with self.assertRaises(EmptyDocumentError):
                extract_document_text(str(temp_path), '.txt')
        finally:
            temp_path.unlink(missing_ok=True)

    def test_document_version_supports_pdf_input_when_available(self):
        try:
            import fitz
        except ImportError:
            self.skipTest('PyMuPDF not available in this environment')

        temp_path = Path('D:/Venus_M2/Venus_project/test_data_demo/demo.pdf')
        temp_path.parent.mkdir(parents=True, exist_ok=True)
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), 'DEMO - page une\nDEMO - page deux')
        doc.save(temp_path)
        doc.close()

        try:
            pages = extract_document_text(str(temp_path), '.pdf')
            self.assertTrue(pages)
            self.assertTrue(all('DEMO' in item['text'] for item in pages))
        finally:
            temp_path.unlink(missing_ok=True)

    def test_docx_extraction_works_with_demo_document(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest('python-docx not available in this environment')

        temp_path = Path('D:/Venus_M2/Venus_project/test_data_demo/demo.docx')
        temp_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_heading('Section Demo', level=1)
        document.add_paragraph('Bonjour le monde. Ceci est un document de démonstration.')
        document.save(str(temp_path))

        try:
            pages = extract_document_text(str(temp_path), '.docx')
            self.assertTrue(pages)
            self.assertTrue(any('Bonjour le monde' in item['text'] for item in pages))
        finally:
            temp_path.unlink(missing_ok=True)

    def test_invalid_file_type_raises(self):
        with self.assertRaises(UnsupportedFileTypeError):
            extract_document_text('demo.unknown', '.unknown')

    def test_missing_file_raises(self):
        with self.assertRaises(MissingDocumentFileError):
            ingest_document_version(
                DocumentVersion(
                    document=Document.objects.create(
                        title='DEMO - Missing file',
                        category=self.category,
                        created_by=self.user,
                    ),
                    version_number=1,
                    file='missing.txt',
                    original_filename='missing.txt',
                    uploaded_by=self.user,
                )
            )

    def test_invalid_pdf_raises_invalid_document_error(self):
        document = Document.objects.create(
            title='DEMO - Invalid PDF',
            category=self.category,
            created_by=self.user,
        )
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('invalid.pdf', b'not a pdf', content_type='application/pdf'),
            original_filename='invalid.pdf',
            uploaded_by=self.user,
        )

        from .services import InvalidDocumentError

        with self.assertRaises(InvalidDocumentError):
            ingest_document_version(version)

    def test_ingestion_is_atomic_when_chunk_creation_fails(self):
        document = Document.objects.create(
            title='DEMO - Atomic ingestion',
            category=self.category,
            created_by=self.user,
        )
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            file=SimpleUploadedFile('atomic.txt', b'new content ' * 20, content_type='text/plain'),
            original_filename='atomic.txt',
            uploaded_by=self.user,
        )
        DocumentChunk.objects.create(
            document_version=version,
            chunk_index=0,
            content='previous content',
        )

        with patch(
            'knowledge.services.DocumentChunk.objects.create',
            side_effect=IntegrityError('simulated chunk failure'),
        ):
            with self.assertRaises(IntegrityError):
                ingest_document_version(version, chunk_size=3, overlap=1)

        self.assertEqual(
            list(version.chunks.values_list('content', flat=True)),
            ['previous content'],
        )
